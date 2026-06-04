from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

BG    = RGBColor(0x18, 0x18, 0x18)
CARD  = RGBColor(0x26, 0x26, 0x26)
LINE  = RGBColor(0x40, 0x40, 0x40)
OR    = RGBColor(0xF2, 0x6A, 0x2F)
GOLD  = RGBColor(0xF4, 0xC4, 0x30)
MINT  = RGBColor(0x2E, 0xCC, 0x9A)
BLUE  = RGBColor(0x3A, 0x86, 0xFF)
WHITE = RGBColor(0xF2, 0xF2, 0xF2)
LGRAY = RGBColor(0x88, 0x88, 0x88)
RED   = RGBColor(0xE0, 0x3E, 0x3E)

W = prs.slide_width
H = prs.slide_height

def bg(slide):
    s = slide.shapes.add_shape(1, 0, 0, W, H)
    s.fill.solid(); s.fill.fore_color.rgb = BG
    s.line.fill.background()

def rect(slide, l, t, w, h, fill, line_col=None, lw=1.0):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line_col: s.line.color.rgb = line_col; s.line.width = Pt(lw)
    else: s.line.fill.background()
    return s

def tx(slide, text, l, t, w, h,
        size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    r  = p.add_run(); r.text = text
    r.font.size=Pt(size); r.font.bold=bold
    r.font.color.rgb=color; r.font.italic=italic
    return tb

def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def bline(slide, color=OR):
    rect(slide, 0, H - Inches(0.06), W, Inches(0.06), color)

def hdr(slide, text, color=WHITE, sub_color=OR):
    rect(slide, 0, 0, W, Inches(1.05), CARD)
    tx(slide, text, Inches(0.7), Inches(0.2), W, Inches(0.65),
       size=30, bold=True, color=color)
    rect(slide, Inches(0.7), Inches(0.9), Inches(len(text)*0.19+0.2), Inches(0.04), sub_color)

def pg(slide, n, total=9):
    tx(slide, f"{n}  /  {total}", W-Inches(1.4), H-Inches(0.42),
       Inches(1.3), Inches(0.32), size=11, color=LGRAY, align=PP_ALIGN.RIGHT)

def tag(slide, text, x, y, color=OR, text_color=None):
    tc = text_color or BG
    w = Inches(len(text)*0.14 + 0.45)
    rect(slide, x, y, w, Inches(0.35), color)
    tx(slide, text, x+Inches(0.1), y+Inches(0.03), w, Inches(0.32),
       size=12, bold=True, color=tc)
    return w

# ══════════════════════════════════════════════════════════════
# 1 — 타이틀
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl, OR)
rect(sl, 0, 0, Inches(0.55), H, OR)
rect(sl, W-Inches(3.2), 0, Inches(3.2), H, CARD)
rect(sl, W-Inches(3.2), 0, Inches(0.04), H, OR)

tx(sl, "VicPinky", Inches(0.95), Inches(1.1), Inches(8), Inches(2.1),
   size=72, bold=True, color=WHITE)
tx(sl, "자율주행 로봇 개발", Inches(0.95), Inches(3.2), Inches(8), Inches(0.7),
   size=30, color=LGRAY)
rect(sl, Inches(0.95), Inches(4.1), Inches(2.8), Inches(0.04), OR)
tx(sl, "중간 발표  ·  2026. 05. 21", Inches(0.95), Inches(4.3), Inches(6), Inches(0.5),
   size=16, color=LGRAY, italic=True)

for i, t in enumerate(["자율주행\nNAV", "사람추종\nFOLLOW", "웹 UI\n전환"]):
    tx(sl, t, W-Inches(2.8), Inches(1.2)+i*Inches(1.8), Inches(2.5), Inches(1.5),
       size=20, bold=True, color=LINE, align=PP_ALIGN.CENTER)

note(sl, '"안녕하세요. 저는 공장 같은 실내 환경에서 두 가지 모드로 동작하는 자율주행 로봇을 만들어봤습니다.\n자율주행 모드와 사람 추종 모드, 그리고 버튼 하나로 전환하는 기능까지 포함합니다."')
pg(sl, 1)

# ══════════════════════════════════════════════════════════════
# 2 — 두 가지 모드
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl)
hdr(sl, "두 가지 모드로 작동해요")

# NAV 모드 카드
rect(sl, Inches(0.5), Inches(1.2), Inches(5.9), Inches(5.7), CARD, MINT, 2)
rect(sl, Inches(0.5), Inches(1.2), Inches(5.9), Inches(0.55), MINT)
tx(sl, "NAV  모드", Inches(0.5), Inches(1.2), Inches(5.9), Inches(0.55),
   size=20, bold=True, color=BG, align=PP_ALIGN.CENTER)
tx(sl, "🗺️", Inches(0.5), Inches(1.9), Inches(5.9), Inches(0.9),
   size=44, align=PP_ALIGN.CENTER)
tx(sl, "목적지를 찍으면\n스스로 찾아가기",
   Inches(0.7), Inches(2.9), Inches(5.5), Inches(1.2),
   size=24, bold=True, color=WHITE)
for line in ["• 여러 경유지 순서대로 이동",
             "• YOLO로 사람/물체 장애물 처리",
             "• 막히면 자동 복구 (후진→대기→회전)"]:
    y_off = ["• 여러 경유지 순서대로 이동",
             "• YOLO로 사람/물체 장애물 처리",
             "• 막히면 자동 복구 (후진→대기→회전)"].index(line)
    tx(sl, line, Inches(0.7), Inches(4.25)+y_off*Inches(0.6),
       Inches(5.5), Inches(0.55), size=15, color=LGRAY)

# 가운데 전환 표시
rect(sl, Inches(6.5), Inches(3.3), Inches(0.35), Inches(1.0),
     RGBColor(0x35,0x35,0x35))
tx(sl, "⇄", Inches(6.3), Inches(3.5), Inches(0.75), Inches(0.6),
   size=28, bold=True, color=OR, align=PP_ALIGN.CENTER)
tx(sl, "웹 UI\n버튼", Inches(6.25), Inches(4.15), Inches(0.85), Inches(0.7),
   size=11, color=OR, align=PP_ALIGN.CENTER)

# FOLLOW 모드 카드
rect(sl, Inches(7.15), Inches(1.2), Inches(5.9), Inches(5.7), CARD, OR, 2)
rect(sl, Inches(7.15), Inches(1.2), Inches(5.9), Inches(0.55), OR)
tx(sl, "FOLLOW  모드", Inches(7.15), Inches(1.2), Inches(5.9), Inches(0.55),
   size=20, bold=True, color=BG, align=PP_ALIGN.CENTER)
tx(sl, "🧍", Inches(7.15), Inches(1.9), Inches(5.9), Inches(0.9),
   size=44, align=PP_ALIGN.CENTER)
tx(sl, "앞에 있는 사람을\n자동으로 따라가기",
   Inches(7.35), Inches(2.9), Inches(5.5), Inches(1.2),
   size=24, bold=True, color=WHITE)
for i, line in enumerate(["• 카메라 + YOLO로 사람 인식",
                           "• 거리·방향 계산 → cmd_vel 직접 제어",
                           "• 전환 시 Nav2 목표 자동 취소"]):
    tx(sl, line, Inches(7.35), Inches(4.25)+i*Inches(0.6),
       Inches(5.5), Inches(0.55), size=15, color=LGRAY)

note(sl, '"로봇에는 두 가지 모드가 있습니다.\nNAV 모드는 목적지를 지정하면 자율주행하는 모드고,\nFOLLOW 모드는 카메라로 사람을 인식해서 따라가는 모드입니다.\n웹 브라우저의 버튼 하나로 언제든 전환할 수 있습니다."')
pg(sl, 2)

# ══════════════════════════════════════════════════════════════
# 3 — 시스템 구성
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl)
hdr(sl, "어떻게 작동하나요?")

# 공통 센서
rect(sl, Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.0), RGBColor(0x22,0x22,0x22), LINE, 1)
tx(sl, "센서", Inches(0.7), Inches(1.25), Inches(1.0), Inches(0.4),
   size=13, bold=True, color=LGRAY)
for i, (icon, label) in enumerate([("📡","LiDAR\n(거리 측정)"),
                                     ("📷","카메라\n(영상)"),
                                     ("⚙️","오도메트리\n(바퀴 이동량)")]):
    x = Inches(2.0) + i*Inches(3.5)
    tx(sl, icon,  x, Inches(1.22), Inches(0.6), Inches(0.5), size=22, align=PP_ALIGN.CENTER)
    tx(sl, label, x+Inches(0.6), Inches(1.22), Inches(2.5), Inches(0.9), size=13, color=WHITE)

# NAV 경로
rect(sl, Inches(0.5), Inches(2.55), Inches(5.8), Inches(4.3), RGBColor(0x18,0x22,0x1E), MINT, 1.5)
tag(sl, "NAV 모드", Inches(0.65), Inches(2.65), MINT, BG)
steps_nav = [("📍","AMCL\n위치 추정"), ("🛣️","경로\n계산"), ("🤖","DWB\n로컬 플래너")]
for i, (icon, label) in enumerate(steps_nav):
    x = Inches(0.8) + i*Inches(1.75)
    rect(sl, x, Inches(3.1), Inches(1.55), Inches(2.9), CARD, MINT, 1)
    tx(sl, icon,  x, Inches(3.2),  Inches(1.55), Inches(0.7), size=28, align=PP_ALIGN.CENTER)
    tx(sl, label, x, Inches(4.05), Inches(1.55), Inches(0.9), size=13, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 2:
        tx(sl, "→", x+Inches(1.55), Inches(3.9), Inches(0.2), Inches(0.5),
           size=16, color=LGRAY, align=PP_ALIGN.CENTER)

# FOLLOW 경로
rect(sl, Inches(6.85), Inches(2.55), Inches(5.95), Inches(4.3), RGBColor(0x22,0x1A,0x12), OR, 1.5)
tag(sl, "FOLLOW 모드", Inches(7.0), Inches(2.65), OR, BG)
steps_fol = [("🎯","YOLO\n사람 인식"), ("📐","거리/방향\n계산"), ("🕹️","cmd_vel\n직접 전송")]
for i, (icon, label) in enumerate(steps_fol):
    x = Inches(7.05) + i*Inches(1.85)
    rect(sl, x, Inches(3.1), Inches(1.6), Inches(2.9), CARD, OR, 1)
    tx(sl, icon,  x, Inches(3.2),  Inches(1.6), Inches(0.7), size=28, align=PP_ALIGN.CENTER)
    tx(sl, label, x, Inches(4.05), Inches(1.6), Inches(0.9), size=13, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 2:
        tx(sl, "→", x+Inches(1.6), Inches(3.9), Inches(0.25), Inches(0.5),
           size=16, color=LGRAY, align=PP_ALIGN.CENTER)

# 공통 출력
rect(sl, Inches(4.3), Inches(5.9), Inches(4.7), Inches(0.9), RGBColor(0x28,0x28,0x28), LINE, 1)
tx(sl, "→  cmd_vel  →  로봇 바퀴 구동",
   Inches(4.3), Inches(5.95), Inches(4.7), Inches(0.7),
   size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

note(sl, '"공통적으로 LiDAR, 카메라, 오도메트리 센서를 씁니다.\nNAV 모드에서는 AMCL로 위치를 추정하고 경로를 계산해서 DWB가 바퀴를 제어합니다.\nFOLLOW 모드에서는 YOLO로 사람을 인식하고 직접 속도 명령을 보냅니다.\n두 모드 모두 최종적으로는 cmd_vel 토픽으로 바퀴를 움직입니다."')
pg(sl, 3)

# ══════════════════════════════════════════════════════════════
# 4 — 목표
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl)
hdr(sl, "우리가 하고 싶었던 것")

for i, (col, num, title, sub) in enumerate([
    (MINT, "01", "버튼 하나로\n모드 전환", "NAV ↔ FOLLOW\n웹 UI에서 즉시 전환"),
    (OR,   "02", "장애물 있으면\n알아서 처리", "사람/물체 인식 후\n우회하거나 대기"),
    (GOLD, "03", "좁은 문도\n막힘 없이 통과", "80cm 문, 복도\n경로 중앙으로 정확하게"),
]):
    x = Inches(0.55) + i*Inches(4.2)
    rect(sl, x, Inches(1.3), Inches(3.9), Inches(5.6), CARD, col, 1.5)
    tx(sl, num, x, Inches(1.35), Inches(3.9), Inches(1.1),
       size=60, bold=True, color=RGBColor(0x33,0x33,0x33), align=PP_ALIGN.CENTER)
    rect(sl, x+Inches(0.3), Inches(2.55), Inches(3.3), Inches(0.04), col)
    tx(sl, title, x+Inches(0.2), Inches(2.75), Inches(3.5), Inches(1.5),
       size=26, bold=True, color=WHITE)
    tx(sl, sub, x+Inches(0.2), Inches(4.5), Inches(3.5), Inches(2.2),
       size=16, color=LGRAY)

note(sl, '"목표는 세 가지였습니다.\n첫째, 웹 UI 버튼 하나로 NAV와 FOLLOW 모드를 즉시 전환하는 것.\n둘째, 이동 중 장애물이 있으면 스스로 처리하는 것.\n셋째, 실제 사무실의 좁은 문도 부딪히지 않고 통과하는 것입니다."')
pg(sl, 4)

# ══════════════════════════════════════════════════════════════
# 5 — 문제들 (4개)
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl, RED)
hdr(sl, "근데...  문제가 생겼어요", RED, RED)

probs = [
    ("😶", "로봇이\n꼼짝 안 함",      "속도 명령이 중간에 끊겼음"),
    ("👻", "경로가 벽을\n뚫고 지나감", "미탐색 구역 통과 허용 설정"),
    ("💥", "문에\n부딪힘",            "안전 거리 < 로봇 크기"),
    ("🔄", "모드 전환 후\n위치를 잃음", "AMCL이 자기 위치\n재추정 못 함"),
]
for i, (icon, title, cause) in enumerate(probs):
    x = Inches(0.4) + i*Inches(3.23)
    col = RED if i < 3 else GOLD
    rect(sl, x, Inches(1.2), Inches(3.0), Inches(5.65), CARD, col, 1.5)
    tx(sl, icon, x, Inches(1.3), Inches(3.0), Inches(0.9),
       size=38, align=PP_ALIGN.CENTER)
    tx(sl, title, x+Inches(0.15), Inches(2.35), Inches(2.7), Inches(1.3),
       size=22, bold=True, color=WHITE)
    rect(sl, x+Inches(0.15), Inches(3.75), Inches(2.7), Inches(0.03), LINE)
    tag(sl, "원인", x+Inches(0.15), Inches(3.88), col if i < 3 else GOLD)
    tx(sl, cause, x+Inches(0.15), Inches(4.35), Inches(2.7), Inches(2.3),
       size=14, color=LGRAY)

# 4번째 문제 강조 (아직 해결 중)
tx(sl, "← 현재 해결 중",
   Inches(9.8), Inches(6.65), Inches(3.0), Inches(0.4),
   size=13, color=GOLD, italic=True)

note(sl, '"만들면서 네 가지 문제가 있었습니다.\n처음 세 가지는 해결했고, 마지막 모드 전환 후 위치를 잃는 문제는 현재 해결 방법을 찾고 있습니다.\n원인은 FOLLOW 모드에서 움직인 후 위치 추정 시스템이 현재 위치를 다시 못 찾는 것입니다."')
pg(sl, 5)

# ══════════════════════════════════════════════════════════════
# 6 — 해결
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl, MINT)
hdr(sl, "이렇게 고쳤어요  (3/4)", MINT, MINT)

fixes = [
    ("로봇이 안 움직임",  "속도 명령 중간에 끊김",  "끊긴 연결 스크립트에서 복구",       OR),
    ("경로가 벽 통과",    "미탐색 구역 통과 가능",   "설정 변경 →\n미탐색 구역 차단",     MINT),
    ("문에 부딪힘",       "안전 거리 너무 작음",     "안전 거리 늘리고\n벽 근처 패널티 강화", GOLD),
]
for i, (prob, cause, fix, col) in enumerate(fixes):
    x = Inches(0.5) + i*Inches(4.25)
    rect(sl, x, Inches(1.2), Inches(3.9), Inches(2.15), RGBColor(0x28,0x14,0x14), RED, 1)
    tag(sl, "문제", x+Inches(0.15), Inches(1.32), RED)
    tx(sl, prob,  x+Inches(0.15), Inches(1.75), Inches(3.6), Inches(0.5),
       size=18, bold=True, color=WHITE)
    tx(sl, cause, x+Inches(0.15), Inches(2.3),  Inches(3.6), Inches(0.9),
       size=14, color=LGRAY)

    tx(sl, "↓", x+Inches(1.55), Inches(3.45), Inches(0.8), Inches(0.5),
       size=26, bold=True, color=col, align=PP_ALIGN.CENTER)

    rect(sl, x, Inches(4.0), Inches(3.9), Inches(2.75), RGBColor(0x14,0x24,0x18), col, 1.8)
    tag(sl, "해결", x+Inches(0.15), Inches(4.12), col)
    tx(sl, fix, x+Inches(0.15), Inches(4.58), Inches(3.6), Inches(2.0),
       size=18, bold=True, color=WHITE)

note(sl, '"세 가지는 모두 설정 문제였습니다.\n첫 번째는 속도 명령이 중간에 끊겼던 거, 스크립트에서 다시 이어줬어요.\n두 번째는 미탐색 구역을 통과 가능으로 설정해뒀던 거, 옵션 하나 끄니까 해결됐습니다.\n세 번째 문 충돌은 안전 거리를 로봇 크기에 맞게 늘리고 벽 근처 패널티를 강화해서 해결했습니다."')
pg(sl, 6)

# ══════════════════════════════════════════════════════════════
# 7 — AMCL 문제 상세
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl, GOLD)
hdr(sl, "현재 과제 — 모드 전환 후 위치 재추정", GOLD, GOLD)

# 왼쪽: 문제 설명
rect(sl, Inches(0.5), Inches(1.2), Inches(5.8), Inches(5.65), CARD, GOLD, 1.5)
tx(sl, "무슨 일이 일어나나요?", Inches(0.7), Inches(1.35), Inches(5.4), Inches(0.5),
   size=18, bold=True, color=GOLD)

flow = [
    (OR,   "①  FOLLOW 모드로 전환"),
    (OR,   "②  로봇이 사람 따라 이동"),
    (RED,  "③  Nav2가 꺼진 상태라\n    위치 추정이 흐트러짐"),
    (RED,  "④  NAV 모드로 돌아오면\n    로봇이 자기 위치를 모름"),
    (GOLD, "⑤  경로 계산 실패 or\n    엉뚱한 방향으로 출발"),
]
for i, (col, text) in enumerate(flow):
    y = Inches(2.0) + i*Inches(0.9)
    rect(sl, Inches(0.7), y, Inches(0.06), Inches(0.7), col)
    tx(sl, text, Inches(0.9), y+Inches(0.05), Inches(5.2), Inches(0.75),
       size=14, color=WHITE if col != GOLD else GOLD)

# 오른쪽: 해결 방향 (검토 중)
rect(sl, Inches(6.8), Inches(1.2), Inches(6.0), Inches(5.65), CARD, LINE, 1)
tx(sl, "해결 방향 검토 중", Inches(7.0), Inches(1.35), Inches(5.6), Inches(0.5),
   size=18, bold=True, color=LGRAY)

options = [
    (MINT, "방법 A",
     "전환 시 /initialpose 발행",
     "AMCL에 '현재 여기 있어' 라고\n강제로 알려주기\n→ 가장 빠른 해결책"),
    (BLUE, "방법 B",
     "AMCL 파라미터 튜닝",
     "파티클 수 늘리고\n업데이트 주기 높여서\nFOLLOW 중에도 추적 유지"),
]
for i, (col, label, title, desc) in enumerate(options):
    y = Inches(2.05) + i*Inches(2.4)
    rect(sl, Inches(7.0), y, Inches(5.6), Inches(2.15), RGBColor(0x20,0x20,0x20), col, 1.2)
    tag(sl, label, Inches(7.15), y+Inches(0.15), col)
    tx(sl, title, Inches(7.15), y+Inches(0.6), Inches(5.2), Inches(0.5),
       size=17, bold=True, color=col)
    tx(sl, desc, Inches(7.15), y+Inches(1.1), Inches(5.2), Inches(1.0),
       size=14, color=LGRAY)

tx(sl, "→ 발표 후 적용 예정",
   Inches(7.0), Inches(6.55), Inches(5.6), Inches(0.4),
   size=14, color=GOLD, italic=True)

note(sl, '"네 번째 문제는 아직 해결 중입니다.\nFOLLOW 모드에서 이동하고 나면, NAV 모드로 돌아왔을 때 로봇이 자기 위치를 제대로 못 찾습니다.\n해결 방향은 두 가지를 검토하고 있는데, 전환할 때 위치를 강제로 알려주는 방법이 제일 간단할 것 같습니다.\n발표 이후에 적용해볼 예정입니다."')
pg(sl, 7)

# ══════════════════════════════════════════════════════════════
# 8 — 현재 결과
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl)
hdr(sl, "지금은 이렇게 돼요")

results = [
    (True,  MINT, "NAV 모드 — 목적지 찍으면 자율주행 정상 동작"),
    (True,  MINT, "FOLLOW 모드 — YOLO로 사람 인식 후 추종"),
    (True,  MINT, "웹 UI 버튼으로 모드 전환"),
    (True,  MINT, "경로가 벽을 피해서 생성"),
    (True,  MINT, "장애물 있으면 피해가거나 기다림"),
    (False, GOLD, "문 통과 — 설정 조정 완료, 실환경 테스트 진행 중"),
    (False, RED,  "모드 전환 후 AMCL 재추정 — 해결 방법 검토 중"),
]
for i, (done, col, text) in enumerate(results):
    y = Inches(1.2) + i*Inches(0.82)
    bg_col = CARD if done else RGBColor(0x24,0x20,0x14) if col==GOLD else RGBColor(0x24,0x14,0x14)
    rect(sl, Inches(0.6), y, Inches(12.1), Inches(0.72), bg_col)
    rect(sl, Inches(0.6), y, Inches(0.07), Inches(0.72), col)
    icon_bg = MINT if done else col
    rect(sl, Inches(0.75), y+Inches(0.13), Inches(0.46), Inches(0.46), icon_bg)
    tx(sl, "✓" if done else ("⚠" if col==GOLD else "✗"),
       Inches(0.75), y+Inches(0.13), Inches(0.46), Inches(0.46),
       size=14, bold=True, color=BG, align=PP_ALIGN.CENTER)
    tx(sl, text, Inches(1.35), y+Inches(0.14), Inches(11.2), Inches(0.46),
       size=17, color=WHITE if done else col, bold=done)

note(sl, '"현재 상태입니다.\n초록 체크는 정상 동작 중이고, 노란 경고는 테스트 중, 빨간 건 현재 해결 중입니다.\n핵심 기능인 자율주행, 사람 추종, 모드 전환은 모두 동작합니다."')
pg(sl, 8)

# ══════════════════════════════════════════════════════════════
# 9 — 앞으로
# ══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl); bline(sl, OR)
hdr(sl, "앞으로 할 것들")

nexts = [
    (RED,  "🔧", "당장",      "AMCL 재추정 문제 해결\n(전환 시 /initialpose 발행)"),
    (MINT, "📐", "이번 주",   "문 통과 실환경 최종 테스트\n파라미터 확정"),
    (GOLD, "📷", "다음 달",   "카메라 장애물 정확도 향상\n사람 외 물체 인식 개선"),
    (BLUE, "🎙️", "그 이후",  "음성 명령 목적지 설정\n스마트폰 원격 제어"),
]
for i, (col, icon, period, body) in enumerate(nexts):
    x = Inches(0.4) + i*Inches(3.25)
    rect(sl, x, Inches(1.2), Inches(3.0), Inches(5.65), CARD, col, 1.5)
    rect(sl, x, Inches(1.2), Inches(3.0), Inches(0.48), col)
    tx(sl, period, x, Inches(1.2), Inches(3.0), Inches(0.48),
       size=14, bold=True, color=BG, align=PP_ALIGN.CENTER)
    tx(sl, icon, x, Inches(1.82), Inches(3.0), Inches(0.85),
       size=38, align=PP_ALIGN.CENTER)
    tx(sl, body, x+Inches(0.15), Inches(2.85), Inches(2.7), Inches(3.8),
       size=16, color=WHITE)

rect(sl, Inches(0.4), Inches(6.82), Inches(12.5), Inches(0.04), LINE)
tx(sl, "최종 목표 :  작업자와 함께 공장을 자유롭게 돌아다니는 서비스 로봇",
   Inches(0.4), Inches(6.92), W-Inches(0.8), Inches(0.45),
   size=16, color=LGRAY, italic=True, align=PP_ALIGN.CENTER)

note(sl, '"가장 급한 건 AMCL 위치 재추정 문제 해결입니다. 이번 주 안에 잡겠습니다.\n그 다음 문 통과 마무리하고, 이후에 카메라 정확도를 높이고 음성 명령까지 붙이는 게 목표입니다.\n최종적으로는 공장에서 작업자와 자연스럽게 함께 움직이는 로봇을 만드는 것입니다.\n이상입니다. 감사합니다."')
pg(sl, 9)

out = "/home/hong/dev_ws/vic_pinky/VicPinky_중간발표_20260521.pptx"
prs.save(out)
print(f"저장: {out}")
